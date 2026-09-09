from pathlib import Path
from document_rag_eval.core import ParsedBlock
import json

def parse_file(path):
    p=Path(path); did=p.stem
    if p.suffix.lower()=='.pdf':
        from pypdf import PdfReader
        return [ParsedBlock(did,str(p),page.extract_text() or '',{'page':i+1},'page') for i,page in enumerate(PdfReader(str(p)).pages)]
    if p.suffix.lower()=='.docx':
        from docx import Document
        d=Document(str(p)); out=[]
        for i,x in enumerate(d.paragraphs):
            if x.text.strip(): out.append(ParsedBlock(did,str(p),x.text,{'paragraph_index':i},'paragraph'))
        for ti,t in enumerate(d.tables):
            for ri,row in enumerate(t.rows):
                for ci,c in enumerate(row.cells): out.append(ParsedBlock(did,str(p),c.text,{'table_index':ti,'row_index':ri,'cell_index':ci},'table_cell'))
        return out
    if p.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(str(p),read_only=True,data_only=True); out=[]
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None: out.append(ParsedBlock(did,str(p),str(c.value),{'sheet':ws.title,'row':c.row,'column':c.column,'cell':c.coordinate},'cell'))
        return out
    raise ValueError(json.dumps({'error':'unsupported_format','path':str(p)}))
